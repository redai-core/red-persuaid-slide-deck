import tempfile
import unittest
from pathlib import Path
from pptx import Presentation

from engine.deckcraft.tokens import CANVAS_WIDTH, CANVAS_HEIGHT
from engine.deckcraft.builder import DeckCraftBuilder, compile_deck


class TestDeckCraft(unittest.TestCase):
    def test_dimensions(self):
        self.assertAlmostEqual(CANVAS_WIDTH.inches, 20.0, places=2)
        self.assertAlmostEqual(CANVAS_HEIGHT.inches, 11.25, places=2)

    def test_build_full_21_slides(self):
        brand_data = {
            "brand": "Electrum",
            "category": "motor listrik",
            "competitors": "Alva, Smoot, Polytron, VinFast, Gesits",
            "domain": "electrum.id",
            "year": 2026,
            "hero_stat": {
                "share_of_voice_pct": 9.3,
                "average_rank": "#4",
                "sentiment_score": "63/100",
                "win_rate_pct": "18%",
            },
        }

        builder = DeckCraftBuilder(brand_data)
        with tempfile.TemporaryDirectory() as tmpdir:
            out_file = Path(tmpdir) / "test_deck.pptx"
            res = builder.build(str(out_file))
            self.assertTrue(res.exists())

            # Load generated presentation and verify integrity
            prs = Presentation(str(res))
            self.assertEqual(len(prs.slides), 21)
            self.assertAlmostEqual(prs.slide_width.inches, 20.0, places=2)
            self.assertAlmostEqual(prs.slide_height.inches, 11.25, places=2)

            for i, slide in enumerate(prs.slides):
                self.assertGreater(len(slide.shapes), 0, f"Slide {i+1} has no shapes")

    def test_compile_deck_with_variable_competitors(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            # Test with 3 competitors
            out_file = compile_deck(
                brand="Siloam Hospitals",
                category="hospital network",
                competitors="Mayapada, Pondok Indah, Mitra Keluarga",
                out_dir=tmpdir,
                year=2026,
            )
            self.assertTrue(out_file.exists())
            prs = Presentation(str(out_file))
            self.assertEqual(len(prs.slides), 21)

    def test_mcp_handle_persuaid_generate_deck(self):
        from engine.mcp_server import handle_persuaid_generate_deck
        with tempfile.TemporaryDirectory() as tmpdir:
            res = handle_persuaid_generate_deck({
                "brand": "Auto2000",
                "category": "authorized Toyota dealer",
                "competitors": "Plaza Toyota, Astrido Toyota, Tunas Toyota",
                "domain": "auto2000.co.id",
                "out_dir": tmpdir,
            })
            self.assertEqual(res["status"], "success")
            self.assertEqual(res["total_slides"], 21)
            self.assertTrue(Path(res["deck_path"]).exists())
            prs = Presentation(res["deck_path"])
            self.assertEqual(len(prs.slides), 21)


if __name__ == "__main__":
    unittest.main()
