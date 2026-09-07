#!/usr/bin/env python3
"""
PersuAId DeckCraft Act II: The Brand Diagnosis (Slides 6 to 11)
6. Section Divider: BRAND REPORT
7. Core Finding Hero Metric Cards + Observation + Cyan Takeaway Banner
8. Section Divider: UNDERSTANDING CONSUMER INTENT
9. The Key Gap: 70/30 Dynamic Competitor Bars + Driver Cards
10. Coverage by Engine: 3 Circular Dials + Dual Diagnosis Cards
11. AI Positioning: 2x2 Verbatim Quotation Cards with “ Glyphs
"""

from typing import Dict, Any, List
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from engine.deckcraft.tokens import (
    CANVAS_WIDTH,
    CANVAS_HEIGHT,
    MARGIN_X,
    MARGIN_Y,
    CONTENT_WIDTH,
    COLOR_OBSIDIAN,
    COLOR_NAVY,
    COLOR_CYAN,
    COLOR_CYAN_DEEP,
    COLOR_CARD_SLATE,
    COLOR_CARD_BORDER,
    COLOR_CARD_MINT,
    COLOR_WHITE,
    COLOR_OFF_WHITE,
    COLOR_COOL_GRAY,
    COLOR_MUTED_SLATE,
    COLOR_DARK_TEXT,
    COLOR_AMBER_WARN,
    FONT_PRIMARY,
    SIZE_SECTION_TITLE,
    SIZE_CARD_HEADER,
    SIZE_BODY,
    SIZE_HERO_STAT,
)
from engine.deckcraft.primitives import (
    set_slide_background,
    add_header,
    add_footer,
    add_card,
    add_stat_card,
    add_takeaway_banner,
)


def build_slide_06_divider_brand(slide, data: Dict[str, Any]):
    """Slide 6: Section Divider: BRAND REPORT."""
    set_slide_background(slide, COLOR_OBSIDIAN)
    brand = data.get("brand", "Your Brand")

    box = slide.shapes.add_textbox(MARGIN_X, Inches(4.5), CONTENT_WIDTH, Inches(2.5))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    p_k = tf.paragraphs[0]
    p_k.text = "DIAGNOSIS & BENCHMARKS"
    p_k.font.name = FONT_PRIMARY
    p_k.font.size = Pt(16)
    p_k.font.bold = True
    p_k.font.color.rgb = COLOR_CYAN
    p_k.alignment = PP_ALIGN.CENTER

    p_t = tf.add_paragraph()
    p_t.text = "BRAND REPORT"
    p_t.font.name = FONT_PRIMARY
    p_t.font.size = SIZE_SECTION_TITLE
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_WHITE
    p_t.alignment = PP_ALIGN.CENTER

    p_s = tf.add_paragraph()
    p_s.text = f"Audit results and search engine baseline for {brand}"
    p_s.font.name = FONT_PRIMARY
    p_s.font.size = Pt(18)
    p_s.font.color.rgb = COLOR_MUTED_SLATE
    p_s.alignment = PP_ALIGN.CENTER

    add_footer(slide, brand)


def build_slide_07_hero_diagnosis(slide, data: Dict[str, Any]):
    """Slide 7: Core Finding Hero Metric Cards + Observation + Cyan Takeaway Banner."""
    set_slide_background(slide, COLOR_OBSIDIAN)
    brand = data.get("brand", "Your Brand")

    hero = data.get("hero_stat", {})
    sov_pct = hero.get("share_of_voice_pct", 9.3)
    avg_rank = hero.get("average_rank", "#4")
    sentiment = hero.get("sentiment_score", "63/100")
    win_rate = hero.get("win_rate_pct", "18%")

    add_header(
        slide=slide,
        kicker="CORE FINDING",
        title="Visible but underpowered — the AI gap is real",
        subtitle=None,
    )

    # 4 Hero Metric Cards Row
    card_w = (CONTENT_WIDTH - Inches(0.9)) / 4  # 4.175 inches
    card_h = Inches(2.2)
    card_y = Inches(2.4)

    # Card 1: Share of Voice
    add_stat_card(
        slide=slide,
        x=MARGIN_X,
        y=card_y,
        w=card_w,
        h=card_h,
        value=f"{sov_pct}%",
        label=f"of AI answers mention {brand}",
        context="Coverage across 225 responses",
    )

    # Card 2: Rank Position
    add_stat_card(
        slide=slide,
        x=MARGIN_X + (card_w + Inches(0.3)),
        y=card_y,
        w=card_w,
        h=card_h,
        value=str(avg_rank),
        label="among category brands",
        context="Behind top 3 competitors",
    )

    # Card 3: Sentiment Score
    add_stat_card(
        slide=slide,
        x=MARGIN_X + (card_w + Inches(0.3)) * 2,
        y=card_y,
        w=card_w,
        h=card_h,
        value=str(sentiment),
        label="average sentiment when named",
        context="Positive but below competitor avg",
    )

    # Card 4: Win Rate (#1 Recommendation)
    add_stat_card(
        slide=slide,
        x=MARGIN_X + (card_w + Inches(0.3)) * 3,
        y=card_y,
        w=card_w,
        h=card_h,
        value=str(win_rate),
        label=f"of mentions place {brand} #1",
        context="On niche/feature queries specifically",
    )

    # Middle Observation Container
    obs_y = Inches(4.9)
    obs_h = Inches(1.8)
    add_card(slide, MARGIN_X, obs_y, CONTENT_WIDTH, obs_h, fill_color=COLOR_CARD_SLATE, border_color=COLOR_CARD_BORDER)

    tb_obs = slide.shapes.add_textbox(MARGIN_X + Inches(0.6), obs_y + Inches(0.25), CONTENT_WIDTH - Inches(1.2), obs_h - Inches(0.5))
    tf_obs = tb_obs.text_frame
    tf_obs.word_wrap = True
    tf_obs.margin_left = tf_obs.margin_right = tf_obs.margin_top = tf_obs.margin_bottom = 0

    obs_lines = [
        f"Unique market advantages exist, but AI rarely highlights them.",
        f"Instead, AI models consistently default to competitors for general purchase queries.",
        f"The gap is visibility and citability, not product credibility.",
    ]
    for i, line in enumerate(obs_lines):
        p = tf_obs.paragraphs[0] if i == 0 else tf_obs.add_paragraph()
        p.text = line
        p.font.name = FONT_PRIMARY
        p.font.size = Pt(17)
        p.font.color.rgb = COLOR_WHITE
        p.alignment = PP_ALIGN.CENTER

    # Signature Solid Cyan Takeaway Banner
    add_takeaway_banner(
        slide=slide,
        points=[
            "The gap is visibility, not credibility.",
            "Strong product differentiators are ignored by AI because they lack editorial citations.",
            "AI is simply defaulting to the brands with the loudest web presence.",
        ],
        y=Inches(7.0),
        h=Inches(1.8),
    )

    add_footer(slide, brand)


def build_slide_08_divider_intent(slide, data: Dict[str, Any]):
    """Slide 8: Section Divider: UNDERSTANDING CONSUMER INTENT."""
    set_slide_background(slide, COLOR_OBSIDIAN)
    brand = data.get("brand", "Your Brand")

    box = slide.shapes.add_textbox(MARGIN_X, Inches(4.5), CONTENT_WIDTH, Inches(2.5))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    p_k = tf.paragraphs[0]
    p_k.text = "CATEGORY SEARCH DYNAMICS"
    p_k.font.name = FONT_PRIMARY
    p_k.font.size = Pt(16)
    p_k.font.bold = True
    p_k.font.color.rgb = COLOR_CYAN
    p_k.alignment = PP_ALIGN.CENTER

    p_t = tf.add_paragraph()
    p_t.text = "UNDERSTANDING CONSUMER INTENT"
    p_t.font.name = FONT_PRIMARY
    p_t.font.size = SIZE_SECTION_TITLE
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_WHITE
    p_t.alignment = PP_ALIGN.CENTER

    p_s = tf.add_paragraph()
    p_s.text = "The queries and prompt journeys that shape AI recommendations"
    p_s.font.name = FONT_PRIMARY
    p_s.font.size = Pt(18)
    p_s.font.color.rgb = COLOR_MUTED_SLATE
    p_s.alignment = PP_ALIGN.CENTER

    add_footer(slide, brand)


def build_slide_09_competitor_gap(slide, data: Dict[str, Any]):
    """Slide 9: The Key Gap: 70/30 Dynamic Competitor Bars + Driver Cards."""
    set_slide_background(slide, COLOR_OBSIDIAN)
    brand = data.get("brand", "Your Brand")

    add_header(
        slide=slide,
        kicker="THE KEY GAP",
        title=f"AI defaults to established players — {brand} barely registers",
        subtitle=None,
    )

    # Competitor Bar Data
    comp_list = data.get("competitor_gap", [])
    if not comp_list:
        comp_list = [
            {"competitor": "Competitor A", "share_of_voice_pct": 31.2},
            {"competitor": "Competitor B", "share_of_voice_pct": 28.4},
            {"competitor": "Competitor C", "share_of_voice_pct": 24.7},
            {"competitor": brand, "share_of_voice_pct": 9.3, "is_client": True},
            {"competitor": "Competitor D", "share_of_voice_pct": 7.4},
        ]

    # Ensure brand is in list with is_client flag
    has_client = any(c.get("is_client") or c.get("competitor", "").lower() == brand.lower() for c in comp_list)
    if not has_client:
        comp_list.append({"competitor": brand, "share_of_voice_pct": 9.3, "is_client": True})

    # Sort descending
    comp_list = sorted(comp_list, key=lambda x: x.get("share_of_voice_pct", 0), reverse=True)
    max_sov = max([c.get("share_of_voice_pct", 10) for c in comp_list] + [35.0])

    # 70/30 Column Layout
    left_w = Inches(12.0)
    right_w = Inches(5.3)
    start_y = Inches(2.4)
    total_bar_h = Inches(5.0)

    row_h = min(Inches(0.65), total_bar_h / len(comp_list))

    # Render Left Column: Horizontal Bars
    for i, comp in enumerate(comp_list):
        row_y = start_y + (i * row_h)
        name = comp.get("competitor", "Brand")
        pct = comp.get("share_of_voice_pct", 0)
        is_client = comp.get("is_client") or name.lower() == brand.lower()

        # Brand Name Label (Left)
        tb_lbl = slide.shapes.add_textbox(MARGIN_X, row_y, Inches(2.2), row_h)
        tf_lbl = tb_lbl.text_frame
        tf_lbl.margin_left = tf_lbl.margin_right = tf_lbl.margin_top = tf_lbl.margin_bottom = 0
        p_lbl = tf_lbl.paragraphs[0]
        p_lbl.text = name
        p_lbl.font.name = FONT_PRIMARY
        p_lbl.font.size = Pt(15)
        p_lbl.font.bold = is_client
        p_lbl.font.color.rgb = COLOR_CYAN if is_client else COLOR_WHITE

        # Bar Fill
        max_bar_len = Inches(8.0)
        bar_len = max(Inches(0.6), max_bar_len * (pct / max_sov))
        bar_x = MARGIN_X + Inches(2.4)

        bar_fill = COLOR_CYAN if is_client else COLOR_CARD_BORDER
        bar_line = COLOR_CYAN if is_client else None
        add_card(slide, bar_x, row_y + Inches(0.08), bar_len, row_h - Inches(0.16), fill_color=bar_fill, border_color=bar_line)

        # Percentage Text
        tb_pct = slide.shapes.add_textbox(bar_x + bar_len + Inches(0.2), row_y, Inches(1.2), row_h)
        tf_pct = tb_pct.text_frame
        tf_pct.margin_left = tf_pct.margin_right = tf_pct.margin_top = tf_pct.margin_bottom = 0
        p_p = tf_pct.paragraphs[0]
        p_p.text = f"{pct}%"
        p_p.font.name = FONT_PRIMARY
        p_p.font.size = Pt(15)
        p_p.font.bold = True
        p_p.font.color.rgb = COLOR_CYAN if is_client else COLOR_WHITE

    # Render Right Column: Why does this happen? Container
    right_x = MARGIN_X + left_w + Inches(0.5)
    add_card(slide, right_x, start_y, right_w, Inches(4.8), fill_color=COLOR_CARD_SLATE, border_color=COLOR_CYAN)

    tb_rw = slide.shapes.add_textbox(right_x + Inches(0.4), start_y + Inches(0.3), right_w - Inches(0.8), Inches(4.2))
    tf_rw = tb_rw.text_frame
    tf_rw.word_wrap = True
    tf_rw.margin_left = tf_rw.margin_right = tf_rw.margin_top = tf_rw.margin_bottom = 0

    p_rh = tf_rw.paragraphs[0]
    p_rh.text = "Why does this happen?"
    p_rh.font.name = FONT_PRIMARY
    p_rh.font.size = Pt(20)
    p_rh.font.bold = True
    p_rh.font.color.rgb = COLOR_WHITE

    drivers = [
        "Competitors dominate industry editorial media",
        "Legacy brand awareness halo in tech & news",
        "Higher density of third-party reviews and forum threads",
        f"⚡ {brand} lacks dedicated consumer-focused coverage",
    ]
    for d in drivers:
        p = tf_rw.add_paragraph()
        p.text = f"\n• {d}"
        p.font.name = FONT_PRIMARY
        p.font.size = Pt(14)
        p.font.color.rgb = COLOR_COOL_GRAY

    # Bottom Takeaway Banner
    add_takeaway_banner(
        slide=slide,
        points=[
            "AI favors brands with stronger content visibility.",
            f"{brand} needs more consumer-focused editorial coverage and citation volume.",
        ],
        y=Inches(7.8),
        h=Inches(1.6),
    )

    add_footer(slide, brand)


def build_slide_10_engine_coverage(slide, data: Dict[str, Any]):
    """Slide 10: Coverage by Engine: 3 Circular Dials + Dual Diagnosis Cards."""
    set_slide_background(slide, COLOR_OBSIDIAN)
    brand = data.get("brand", "Your Brand")

    add_header(
        slide=slide,
        kicker="COVERAGE BY ENGINE",
        title="Consistently low — a supply problem, not a platform problem",
        subtitle=None,
    )

    dial_w = Inches(5.4)
    dial_h = Inches(3.8)
    dial_y = Inches(2.4)
    gap = (CONTENT_WIDTH - (dial_w * 3)) / 2

    # Dial 1: ChatGPT
    d1_x = MARGIN_X
    add_card(slide, d1_x, dial_y, dial_w, dial_h, fill_color=COLOR_CARD_SLATE, border_color=COLOR_CYAN, shape_type=MSO_SHAPE.OVAL)
    tb_d1 = slide.shapes.add_textbox(d1_x + Inches(0.5), dial_y + Inches(0.8), dial_w - Inches(1.0), Inches(2.2))
    tf_d1 = tb_d1.text_frame
    tf_d1.margin_left = tf_d1.margin_right = tf_d1.margin_top = tf_d1.margin_bottom = 0
    p1_v = tf_d1.paragraphs[0]
    p1_v.text = "10.7%"
    p1_v.font.name = FONT_PRIMARY
    p1_v.font.size = SIZE_HERO_STAT
    p1_v.font.bold = True
    p1_v.font.color.rgb = COLOR_CYAN
    p1_v.alignment = PP_ALIGN.CENTER
    p1_e = tf_d1.add_paragraph()
    p1_e.text = "ChatGPT (75 responses)"
    p1_e.font.name = FONT_PRIMARY
    p1_e.font.size = Pt(16)
    p1_e.font.color.rgb = COLOR_WHITE
    p1_e.alignment = PP_ALIGN.CENTER

    # Dial 2: Google AI Overview
    d2_x = MARGIN_X + dial_w + gap
    add_card(slide, d2_x, dial_y, dial_w, dial_h, fill_color=COLOR_CARD_SLATE, border_color=COLOR_CYAN, shape_type=MSO_SHAPE.OVAL)
    tb_d2 = slide.shapes.add_textbox(d2_x + Inches(0.5), dial_y + Inches(0.8), dial_w - Inches(1.0), Inches(2.2))
    tf_d2 = tb_d2.text_frame
    tf_d2.margin_left = tf_d2.margin_right = tf_d2.margin_top = tf_d2.margin_bottom = 0
    p2_v = tf_d2.paragraphs[0]
    p2_v.text = "9.3%"
    p2_v.font.name = FONT_PRIMARY
    p2_v.font.size = SIZE_HERO_STAT
    p2_v.font.bold = True
    p2_v.font.color.rgb = COLOR_CYAN
    p2_v.alignment = PP_ALIGN.CENTER
    p2_e = tf_d2.add_paragraph()
    p2_e.text = "AI Overview (75 responses)"
    p2_e.font.name = FONT_PRIMARY
    p2_e.font.size = Pt(16)
    p2_e.font.color.rgb = COLOR_WHITE
    p2_e.alignment = PP_ALIGN.CENTER

    # Dial 3: Perplexity
    d3_x = MARGIN_X + (dial_w + gap) * 2
    add_card(slide, d3_x, dial_y, dial_w, dial_h, fill_color=COLOR_CARD_SLATE, border_color=COLOR_CYAN, shape_type=MSO_SHAPE.OVAL)
    tb_d3 = slide.shapes.add_textbox(d3_x + Inches(0.5), dial_y + Inches(0.8), dial_w - Inches(1.0), Inches(2.2))
    tf_d3 = tb_d3.text_frame
    tf_d3.margin_left = tf_d3.margin_right = tf_d3.margin_top = tf_d3.margin_bottom = 0
    p3_v = tf_d3.paragraphs[0]
    p3_v.text = "8.0%"
    p3_v.font.name = FONT_PRIMARY
    p3_v.font.size = SIZE_HERO_STAT
    p3_v.font.bold = True
    p3_v.font.color.rgb = COLOR_CYAN
    p3_v.alignment = PP_ALIGN.CENTER
    p3_e = tf_d3.add_paragraph()
    p3_e.text = "Perplexity (75 responses)"
    p3_e.font.name = FONT_PRIMARY
    p3_e.font.size = Pt(16)
    p3_e.font.color.rgb = COLOR_WHITE
    p3_e.alignment = PP_ALIGN.CENTER

    # Bottom Dual Diagnostic Cards
    diag_w = (CONTENT_WIDTH - Inches(0.8)) / 2
    diag_h = Inches(2.4)
    diag_y = Inches(6.8)

    # Left: Warning Card
    add_card(slide, MARGIN_X, diag_y, diag_w, diag_h, fill_color=COLOR_WHITE, border_color=None)
    tb_warn = slide.shapes.add_textbox(MARGIN_X + Inches(0.6), diag_y + Inches(0.4), diag_w - Inches(1.2), diag_h - Inches(0.8))
    tf_w = tb_warn.text_frame
    tf_w.word_wrap = True
    tf_w.margin_left = tf_w.margin_right = tf_w.margin_top = tf_w.margin_bottom = 0
    pw_h = tf_w.paragraphs[0]
    pw_h.text = "⚠  No single engine to fix"
    pw_h.font.name = FONT_PRIMARY
    pw_h.font.size = Pt(18)
    pw_h.font.bold = True
    pw_h.font.color.rgb = COLOR_DARK_TEXT
    pw_b = tf_w.add_paragraph()
    pw_b.text = "\nCoverage is flat at 8–11% everywhere. This is a category-wide content supply problem, not a platform-specific gap."
    pw_b.font.name = FONT_PRIMARY
    pw_b.font.size = Pt(15)
    pw_b.font.color.rgb = COLOR_DARK_TEXT

    # Right: Solution Card
    add_card(slide, MARGIN_X + diag_w + Inches(0.8), diag_y, diag_w, diag_h, fill_color=COLOR_CARD_MINT, border_color=COLOR_CYAN)
    tb_sol = slide.shapes.add_textbox(MARGIN_X + diag_w + Inches(1.4), diag_y + Inches(0.4), diag_w - Inches(1.2), diag_h - Inches(0.8))
    tf_s = tb_sol.text_frame
    tf_s.word_wrap = True
    tf_s.margin_left = tf_s.margin_right = tf_s.margin_top = tf_s.margin_bottom = 0
    ps_h = tf_s.paragraphs[0]
    ps_h.text = "✓  One workstream lifts all three"
    ps_h.font.name = FONT_PRIMARY
    ps_h.font.size = Pt(18)
    ps_h.font.bold = True
    ps_h.font.color.rgb = COLOR_DARK_TEXT
    ps_b = tf_s.add_paragraph()
    ps_b.text = f"\nBecause all engines read the same thin source base, more {brand}-relevant consumer content and forum presence will move all three platforms at once."
    ps_b.font.name = FONT_PRIMARY
    ps_b.font.size = Pt(15)
    ps_b.font.color.rgb = COLOR_DARK_TEXT

    add_footer(slide, brand)


def build_slide_11_verbatim_quotes(slide, data: Dict[str, Any]):
    """Slide 11: AI Positioning: 2x2 Verbatim Quotation Cards with “ Glyphs."""
    set_slide_background(slide, COLOR_OBSIDIAN)
    brand = data.get("brand", "Your Brand")

    add_header(
        slide=slide,
        kicker="POSITIONING",
        title=f"How AI describes {brand}, in its own words",
        subtitle=None,
    )

    quotes = [
        ("A key player backed by institutional investors, practical for daily users and commuters.", "perplexity"),
        ("Strong operational presence, noted for network convenience and enterprise scale.", "Google AI Overview"),
        ("Recognized for infrastructure reliability; frequently associated with fleet logistics.", "ChatGPT"),
        ("Less prominent in direct consumer purchase recommendations; cited mainly for commercial utility.", "ChatGPT"),
    ]

    card_w = (CONTENT_WIDTH - Inches(0.8)) / 2
    card_h = Inches(2.2)
    grid_y = Inches(2.4)

    for i, (q_text, eng) in enumerate(quotes):
        col = i % 2
        row = i // 2
        cx = MARGIN_X + (col * (card_w + Inches(0.8)))
        cy = grid_y + (row * (card_h + Inches(0.4)))

        # White Rounded Card
        add_card(slide, cx, cy, card_w, card_h, fill_color=COLOR_WHITE, border_color=None)

        # Content Box
        tb = slide.shapes.add_textbox(cx + Inches(0.5), cy + Inches(0.3), card_w - Inches(1.0), card_h - Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        # Quote glyph
        p_q = tf.paragraphs[0]
        p_q.text = f"“ {q_text}”"
        p_q.font.name = FONT_PRIMARY
        p_q.font.size = Pt(15)
        p_q.font.italic = True
        p_q.font.color.rgb = COLOR_DARK_TEXT

        # Engine Badge
        p_eng = tf.add_paragraph()
        p_eng.text = f"\n— {eng}"
        p_eng.font.name = FONT_PRIMARY
        p_eng.font.size = Pt(13)
        p_eng.font.bold = True
        p_eng.font.color.rgb = COLOR_CYAN_DEEP

    # Bottom Takeaway Banner
    add_takeaway_banner(
        slide=slide,
        points=[
            f"Keep the commercial/infrastructure credibility.",
            f"Add the direct consumer story to appear in 'best {data.get('category', 'product')}' recommendations.",
        ],
        y=Inches(7.8),
        h=Inches(1.6),
    )

    add_footer(slide, brand)
